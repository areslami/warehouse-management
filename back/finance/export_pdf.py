from decimal import Decimal
from pathlib import Path

from django.conf import settings
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.cell_range import CellRange

from .models import SalesProforma, PurchaseProforma
from .utils import (
    jalali_date,
    load_sales_proforma_template,
    party_name,
    party_national_code,
    rial_words,
    workbook_to_pdf_bytes,
)

FA2EN = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
BOLD = Font(name="B Nazanin", size=14, bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def to_en(value) -> str:
    if value is None:
        return ""
    return str(value).translate(FA2EN)


def clear_region(ws, min_row, max_row, min_col=1, max_col=15):
    for cr in [
        CellRange(str(mr))
        for mr in ws.merged_cells.ranges
        if min_row <= mr.min_row <= max_row
    ]:
        ws.merged_cells.ranges.discard(cr)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            obj = ws._cells.get((row, col))
            if obj is not None and isinstance(obj, MergedCell):
                del ws._cells[(row, col)]


def write_merged(ws, row, start_col, end_col, value, alignment=None, font=None, end_row=None):
    ws.merge_cells(
        start_row=row, start_column=start_col,
        end_row=end_row or row, end_column=end_col,
    )
    cell = ws.cell(row=row, column=start_col)
    cell.value = value
    cell.alignment = alignment or CENTER
    cell.font = font or BOLD
    return cell


def fill_header(ws, proforma, party, kind):
    for col in range(3, 11):
        ws.cell(row=3, column=col).value = None
    write_merged(
        ws, 3, 6, 10,
        "پیش فاکتور فروش" if kind == "sales" else "پیش فاکتور خرید",
    )

    ws["N2"].value = to_en(proforma.serial_number)
    ws["N3"].value = to_en(jalali_date(proforma.date))

    write_merged(
        ws, 5, 1, 15,
        "مشخصات فروشنده" if kind == "sales" else "مشخصات خریدار",
    )
    write_merged(
        ws, 14, 1, 15,
        "مشخصات خریدار" if kind == "sales" else "مشخصات تامین کننده",
    )

    ws["C16"].value = to_en(party_name(party))
    ws["I16"].value = to_en(party_national_code(party))
    ws["M16"].value = to_en(getattr(party, "economic_code", "") or "")
    ws["C18"].value = to_en(getattr(party, "postal_code", "") or "")
    ws["I18"].value = to_en(getattr(party, "phone", "") or "")
    ws["C20"].value = to_en(getattr(party, "address", "") or "")


def fill_lines(ws, proforma):
    subtotal = Decimal(0)
    for idx, line in enumerate(proforma.lines.all()):
        row = 24 + idx
        weight = Decimal(line.weight or 0)
        unit_price = Decimal(line.unit_price or 0)
        tax = Decimal(line.tax or 0)
        discount = Decimal(line.discount or 0)
        line_subtotal = weight * unit_price
        line_total = line_subtotal * (Decimal(1) + tax - discount)

        ws[f"B{row}"].value = to_en(getattr(line.product, "name", "") or "")
        ws[f"F{row}"].value = weight
        ws[f"H{row}"].value = unit_price
        ws[f"K{row}"].value = tax
        ws[f"L{row}"].value = discount
        ws[f"I{row}"].value = line_subtotal
        ws[f"N{row}"].value = line_total

        for col in ("F", "H", "K", "L", "I", "N"):
            ws[f"{col}{row}"].number_format = "0.####"

        subtotal += line_total
    return subtotal


def fill_description(ws, preset):
    desc_start = 29
    right_lines, left_lines = [], []

    if preset and preset.description:
        all_lines = [l.strip() for l in preset.description.splitlines() if l.strip()]
        mid = (len(all_lines) + 1) // 2
        right_lines = all_lines[:mid]
        left_lines = all_lines[mid:]

    max_lines = max(len(right_lines), len(left_lines))
    total_rows = max(max_lines, 4)
    extra = max(total_rows - 4, 0)
    desc_end = desc_start + total_rows - 1
    med = Side(style="thin")
    rtl = Alignment(horizontal="right", vertical="center", readingOrder=2, wrap_text=True)
    small_font = Font(name="B Nazanin", size=11.5)

    clear_region(ws, 29, 36 + extra + 2, max_col=11)
    if extra > 0:
        ws.insert_rows(33, amount=extra)
        clear_region(ws, 33, 36 + extra + 2, max_col=11)

    for i in range(total_rows):
        row = desc_start + i
        ws.row_dimensions[row].height = 30.0 if i < max_lines else 3.0
        for col in range(1, 12):
            b = {}
            if i == 0:
                b["top"] = med
            if i == total_rows - 1:
                b["bottom"] = med
            if col in (1, 3, 8):
                b["left"] = med
            if col in (2, 7, 11):
                b["right"] = med
            ws.cell(row=row, column=col).border = Border(**b)

    write_merged(
        ws, desc_start, 1, 2, "  :توضیحات",
        font=Font(name="B Nazanin", size=14), end_row=desc_end,
    )

    for i in range(total_rows):
        row = desc_start + i

        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        cell = ws.cell(row=row, column=3)
        cell.value = to_en(right_lines[i]) if i < len(right_lines) else ""
        cell.alignment = rtl
        cell.font = small_font

        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=11)
        cell = ws.cell(row=row, column=8)
        cell.value = to_en(left_lines[i]) if i < len(left_lines) else ""
        cell.alignment = rtl
        cell.font = small_font

    for r in range(29, desc_end + 1):
        ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
        ws.merge_cells(start_row=r, start_column=14, end_row=r, end_column=15)

    return extra


def fill_footer(ws, total_payable, extra, kind):
    row = 33 + extra
    sig1 = 35 + extra
    sig2 = 36 + extra
    rtl_nw = Alignment(horizontal="right", vertical="center", readingOrder=2)

    write_merged(ws, row, 1, 7, "مبلغ قابل پرداخت :", rtl_nw)
    write_merged(ws, row, 8, 11, to_en(rial_words(total_payable)))
    write_merged(ws, row, 12, 13, "جمع کل", font=Font(name="B Nazanin", size=14))
    cell = write_merged(ws, row, 14, 15, total_payable)
    cell.number_format = "0.####"

    write_merged(ws, sig1, 4, 5, "فروشنده")
    write_merged(ws, sig1, 12, 13, "خریدار")
    write_merged(ws, sig2, 4, 5, "مهر و امضا")
    write_merged(ws, sig2, 12, 13, "مهر و امضا")

    emza_path = Path(settings.BASE_DIR) / "templates" / "emza.png"
    if emza_path.exists():
        img = XlImage(str(emza_path))
        img.width = 150
        img.height = 60
        anchor = f"D{sig1}" if kind == "sales" else f"L{sig2}"
        ws.add_image(img, anchor)


def render_proforma_pdf(proforma, kind: str) -> bytes:
    if not proforma.pk:
        raise ValueError("Proforma must be saved before export.")

    model = SalesProforma if kind == "sales" else PurchaseProforma
    party_attr = "customer" if kind == "sales" else "supplier"

    proforma = (
        model.objects
        .select_related(party_attr, "export_preset")
        .prefetch_related("lines__product")
        .get(pk=proforma.pk)
    )

    wb, ws = load_sales_proforma_template()
    party = getattr(proforma, party_attr)
    preset = proforma.export_preset

    fill_header(ws, proforma, party, kind)
    subtotal = fill_lines(ws, proforma)

    total_payable = subtotal * (
        Decimal(1)
        - Decimal(proforma.discount or 0)
        + Decimal(proforma.tax or 0)
        + Decimal(proforma.shipping_cost or 0)
        + Decimal(proforma.commission or 0)
        + Decimal(proforma.other_cost or 0)
    )

    for ref, val in [
        ("N26", proforma.shipping_cost),
        ("N27", proforma.commission),
        ("N28", proforma.discount),
        ("N29", proforma.tax),
        ("N30", proforma.other_cost),
    ]:
        ws[ref].value = Decimal(val or 0)
        ws[ref].number_format = "0.####"

    if preset:
        ws["A27"].value = to_en(
            f"حساب بانک {preset.bank_name}:           {preset.account_number}"
        )
        ws["H27"].value = to_en(preset.sheba_number)

    extra = fill_description(ws, preset)
    fill_footer(ws, total_payable, extra, kind)

    last_row = 36 + extra
    for row_num in range(1, last_row + 1):
        rd = ws.row_dimensions[row_num]
        h = rd.height or ws.sheet_format.defaultRowHeight or 15
        if h > 10:
            rd.height = h + 4

    med = Side(style="medium")
    for r in range(1, last_row + 1):
        for c in (1, 15):
            cell = ws.cell(row=r, column=c)
            ob = cell.border
            if c == 1:
                cell.border = Border(left=med, top=ob.top, bottom=ob.bottom, right=ob.right)
            else:
                cell.border = Border(right=med, top=ob.top, bottom=ob.bottom, left=ob.left)
    for c in range(1, 16):
        for r in (1, last_row):
            cell = ws.cell(row=r, column=c)
            ob = cell.border
            if r == 1:
                cell.border = Border(top=med, left=ob.left, right=ob.right, bottom=ob.bottom)
            else:
                cell.border = Border(bottom=med, left=ob.left, right=ob.right, top=ob.top)

    return workbook_to_pdf_bytes(
        wb, ws,
        filename_stem=f"{kind}-proforma-{proforma.pk}",
        print_area=f"A1:O{last_row}",
    )


def render_sales_proforma_pdf(proforma: SalesProforma) -> bytes:
    return render_proforma_pdf(proforma, "sales")


def render_purchase_proforma_pdf(proforma: PurchaseProforma) -> bytes:
    return render_proforma_pdf(proforma, "purchase")
