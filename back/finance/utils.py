import os
import shutil
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path

import jdatetime
from django.conf import settings
from django.utils import timezone
from openpyxl import load_workbook


_ONES = {
    1: "یک",
    2: "دو",
    3: "سه",
    4: "چهار",
    5: "پنج",
    6: "شش",
    7: "هفت",
    8: "هشت",
    9: "نه",
}
_TENS = {
    2: "بیست",
    3: "سی",
    4: "چهل",
    5: "پنجاه",
    6: "شصت",
    7: "هفتاد",
    8: "هشتاد",
    9: "نود",
}
_TEENS = {
    10: "ده",
    11: "یازده",
    12: "دوازده",
    13: "سیزده",
    14: "چهارده",
    15: "پانزده",
    16: "شانزده",
    17: "هفده",
    18: "هجده",
    19: "نوزده",
}
_HUNDREDS = {
    1: "صد",
    2: "دویست",
    3: "سیصد",
    4: "چهارصد",
    5: "پانصد",
    6: "ششصد",
    7: "هفتصد",
    8: "هشتصد",
    9: "نهصد",
}
_SCALES = ("", "هزار", "میلیون", "میلیارد", "تریلیون", "کوادریلیون")


def jalali_date(dt) -> str:
    if not dt:
        return ""
    if timezone.is_aware(dt):
        dt = timezone.localtime(dt)
    return jdatetime.datetime.fromgregorian(datetime=dt).strftime("%Y/%m/%d")


def party_name(party) -> str:
    if not party:
        return ""
    if getattr(party, "customer_type", None) == "corporate":
        return party.company_name or ""
    return party.full_name or ""


def party_national_code(party) -> str:
    if not party:
        return ""
    if getattr(party, "customer_type", None) == "corporate":
        return party.national_id or ""
    return party.personal_code or ""


def _chunk_words(num: int) -> str:
    if num <= 0:
        return ""

    parts: list[str] = []
    h, r = divmod(num, 100)
    if h:
        parts.append(_HUNDREDS[h])
    if r:
        if 10 <= r <= 19:
            parts.append(_TEENS[r])
        else:
            t, o = divmod(r, 10)
            if t:
                parts.append(_TENS[t])
            if o:
                parts.append(_ONES[o])
    return " و ".join(parts)


def number_to_persian_words(value: int) -> str:
    if value == 0:
        return "صفر"
    if value < 0:
        return f"منفی {number_to_persian_words(abs(value))}"

    parts: list[str] = []
    n = value
    scale_idx = 0
    while n and scale_idx < len(_SCALES):
        n, chunk = divmod(n, 1000)
        if chunk:
            chunk_words = _chunk_words(chunk)
            scale = _SCALES[scale_idx]
            parts.append(f"{chunk_words} {scale}".strip())
        scale_idx += 1
    return " و ".join(reversed(parts))


def rial_words(value: Decimal) -> str:
    try:
        n = int(value)
    except Exception:
        n = 0
    return f"{number_to_persian_words(n)} ریال"


def split_two_column_text(text: str, sep: str = "\n---\n") -> tuple[str, str]:
    cleaned = (text or "").strip()
    if not cleaned:
        return "", ""
    if sep in cleaned:
        left, right = cleaned.split(sep, 1)
        return left.strip(), right.strip()
    return cleaned, ""


def resolve_sales_proforma_template_path() -> Path:
    filename = "پیش_فاکتور_ذرت_روس_آرین_طعم_خزر.xlsx"
    base = Path(settings.BASE_DIR) / "templates"
    direct = base / filename
    if direct.exists():
        return direct
    candidates = list(base.glob("پیش_فاکتور*.xlsx"))
    if candidates:
        return candidates[0]
    raise FileNotFoundError(f"Template XLSX not found under {base}")


def load_sales_proforma_template(sheet_names: tuple[str, ...] = ("فرم خام (2)", "فرم خام")):
    wb = load_workbook(resolve_sales_proforma_template_path())
    ws = next((wb[name] for name in sheet_names if name in wb.sheetnames), wb.active)
    return wb, ws


def soffice_cmd() -> str:
    override = (os.environ.get("LIBREOFFICE_PATH") or "").strip()
    if override:
        return override
    for cmd in ("soffice", "libreoffice"):
        found = shutil.which(cmd)
        if found:
            return found
    raise RuntimeError("LibreOffice (libreoffice/soffice) is not installed on the server.")


def workbook_to_pdf_bytes(wb, ws, *, filename_stem: str, print_area: str = "A1:O36") -> bytes:
    for other in list(wb.worksheets):
        if other != ws:
            wb.remove(other)
    wb.active = 0
    ws.print_area = print_area

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        xlsx_path = tmpdir_path / f"{filename_stem}.xlsx"
        wb.save(xlsx_path)

        env = os.environ.copy()
        env.setdefault("HOME", tmpdir)
        subprocess.run(
            [
                soffice_cmd(),
                "--headless",
                "--nologo",
                "--nolockcheck",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                tmpdir,
                str(xlsx_path),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            env=env,
        )

        pdf_path = tmpdir_path / f"{xlsx_path.stem}.pdf"
        if not pdf_path.exists():
            raise RuntimeError("PDF conversion failed.")
        return pdf_path.read_bytes()

